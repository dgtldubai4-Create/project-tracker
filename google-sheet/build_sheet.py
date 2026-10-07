"""Build the Dabur digital tracker workbook (.xlsx) for upload to Google Drive as a Google Sheet."""
import json, glob, sys, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

DB = sys.argv[1]
OUT = sys.argv[2]

def load(c):
    out = {}
    for f in glob.glob(os.path.join(DB, c, "*.json")):
        d = json.load(open(f))
        out[os.path.basename(f)[:-5]] = d
    return out

people = load("people"); markets = load("markets"); projects = load("projects")
settings = json.load(open(os.path.join(DB, "settings", "main.json")))

# ---------- vocab ----------
STREAMS = {"performance": "Performance Marketing", "influencer": "Influencer Marketing", "creators": "Content Creators",
           "social": "Social Media Management", "web": "Website / Tech", "seo": "SEO / GEO", "ai": "AI Projects",
           "bigbets": "Big Bets", "activations": "Digital Activations", "production": "Digital Production", "orm": "ORM",
           "others": "Others", "dashboards": "Dashboards", "sentiment": "Sentiment Analysis"}
STATUS = {"todo": "Not started", "doing": "In progress", "stuck": "Blocked", "done": "Done"}
STATUS_LIST = ["Not started", "In progress", "Blocked", "Done"]
MARKET_LIST = ["MENA", "SSA", "India", "Egypt", "Turkey (Hobby)", "Pakistan", "Multi-market", "Internal"]
MKT = {k: v["name"] for k, v in markets.items()}
AGENCIES = sorted(v["name"] for v in people.values() if v.get("kind") == "partner")
OWNERS = ["Suhail", "Nour", "Sanya", "Khushii", "Team member 4"]
PEOPLE_LIST = OWNERS + ["Manish", "Nishant"]
TYPES = ["Update", "Idea", "Learning", "Expectation", "PPT / deck", "Question for Suhail", "Urgent flag"]
PID = {k: v["name"] for k, v in people.items()}

def d(s):
    try: return dt.date.fromisoformat(s[:10])
    except Exception: return None

rows_by_owner = {o: [] for o in OWNERS}
for p in sorted(projects.values(), key=lambda p: (p.get("due") or "9999", p.get("title", ""))):
    if p.get("archived"): continue
    owner = PID.get(p.get("owner"), "Suhail")
    mk = [MKT.get(m, "") for m in p.get("markets") or []]
    market = mk[0] if len(mk) == 1 else ("Multi-market" if mk else "")
    upd = (p.get("updates") or [])
    latest = (upd[0].get("t") if upd else "") or p.get("nextStep") or ""
    rows_by_owner.setdefault(owner, []).append([
        p.get("title", "").strip(), market, ", ".join(p.get("brands") or []),
        STREAMS.get(p.get("stream"), ""), STATUS.get(p.get("status"), "Not started"),
        d(p.get("due") or ""), latest, p.get("stuckOn") or "",
        ", ".join(PID.get(a, "") for a in (p.get("agencies") or []) if a in PID), "",
        d(p.get("updatedAt") or "")])

# ---------- style ----------
FONT = "Arial"
GREEN = "0E3A2A"; GREEN_SOFT = "DCEFE4"; INK = "12231C"; MUTED = "5A6B62"
title_font = Font(name=FONT, size=15, bold=True, color=INK)
note_font = Font(name=FONT, size=10, italic=True, color=MUTED)
hdr_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=GREEN)
body_font = Font(name=FONT, size=10, color=INK)
sec_font = Font(name=FONT, size=12, bold=True, color=INK)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Border(bottom=Side(style="thin", color="D5DFD9"))
DATE_FMT = "d mmm yyyy"
RED_FILL = PatternFill("solid", fgColor="F8D7D3", bgColor="F8D7D3")
AMBER_FILL = PatternFill("solid", fgColor="FCEBC7", bgColor="FCEBC7")
GREY_FONT = Font(color="8A938E")
ORANGE_FONT = Font(color="B45309", bold=True)
LAST = 1000

def header(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30

def widths(ws, ws_widths):
    for col, w in ws_widths.items(): ws.column_dimensions[col].width = w

def body(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.font = body_font; c.alignment = wrap

wb = Workbook()

# ---------- Start here ----------
st = wb.active; st.title = "Start here"
st.sheet_properties.tabColor = GREEN
st["B2"] = "Dabur International Digital: project tracker"; st["B2"].font = Font(name=FONT, size=18, bold=True, color=INK)
st["B3"] = "Suhail's digital marketing team: MENA and SSA, with support for India, Egypt, Turkey (Hobby) and Pakistan."
st["B3"].font = note_font
lines = [
    ("How to use it", None),
    ("1", "Find your own tab (Suhail, Nour, Sanya, Khushii). Only you and Suhail can edit your tab."),
    ("2", "Add one row per project. Pick Market, Brand, Stream and Status from the dropdowns."),
    ("3", "When something changes, rewrite \"Latest update / next step\" and set \"Last updated\" to today (Ctrl + ; types today's date)."),
    ("4", "If you are stuck or need Suhail, set Status to Blocked and write what you need in \"Blocked on / need from Suhail\"."),
    ("5", "For updates, ideas, learnings, deck links or questions, add a row in the Comments tab. Everyone can write there."),
    ("6", "To comment on any single cell, on any tab, right-click it and choose Comment. Type @ and a name to email that person."),
    ("", ""),
    ("The tabs", None),
    ("Dashboard", "Suhail's overview: counts per person and market, what is overdue or blocked, latest comments. Fills itself in."),
    ("All projects", "Every project from every person's tab in one list. Fills itself in, so don't type here."),
    ("Your tab", "Your projects. This is the only place you add or change a project."),
    ("Comments", "Shared notes for the whole team, including Manish and Nishant. Suhail replies in the last column."),
    ("Lists", "The dropdown values (markets, brands, streams, agencies). Ask Suhail to add a new brand or agency."),
    ("", ""),
    ("Colours", None),
    ("Red row", "Overdue: the due date has passed and the project is not done."),
    ("Amber row", "Blocked: waiting on someone or something."),
    ("Grey text", "Done."),
    ("Orange date", "Last updated more than 14 days ago."),
]
r = 5
for a, b in lines:
    if b is None:
        st.cell(row=r, column=2, value=a).font = sec_font
    else:
        st.cell(row=r, column=2, value=a).font = Font(name=FONT, size=10, bold=True, color=INK)
        c = st.cell(row=r, column=3, value=b); c.font = body_font; c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
widths(st, {"A": 3, "B": 16, "C": 100})
st.sheet_view.showGridLines = False

# ---------- Dashboard (built after All projects so refs are known) ----------
dash = wb.create_sheet("Dashboard"); dash.sheet_properties.tabColor = "1D7752"
allp = wb.create_sheet("All projects"); allp.sheet_properties.tabColor = "2C4078"

# ---------- owner tabs ----------
OWN_HDR = ["Project", "Market", "Brand", "Stream", "Status", "Due date", "Latest update / next step",
           "Blocked on / need from Suhail", "Agency", "Link (brief, deck, folder)", "Last updated"]
OWN_W = {"A": 38, "B": 15, "C": 20, "D": 22, "E": 13, "F": 12, "G": 48, "H": 32, "I": 22, "J": 26, "K": 13}
TAB_COLORS = ["0E3A2A", "B8601A", "0B7D7A", "7A4BA6", "8C6A0E"]

def add_validations(ws, cols):
    """cols: dict of column letter -> (list formula, strict)"""
    for col, (src, strict) in cols.items():
        dv = DataValidation(type="list", formula1=src, allow_blank=True,
                            showErrorMessage=strict, showDropDown=False)
        if strict:
            dv.error = "Pick a value from the list."; dv.errorTitle = "Not on the list"
        ws.add_data_validation(dv); dv.add(f"{col}3:{col}{LAST}")

def add_date_validation(ws, col):
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                        showErrorMessage=True, error="Enter a date, for example 30 Oct 2026.", errorTitle="Date")
    ws.add_data_validation(dv); dv.add(f"{col}3:{col}{LAST}")

def row_colours(ws, last_col, status_col, due_col, upd_col):
    rng = f"A3:{last_col}{LAST}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND(${due_col}3<>"",${due_col}3<TODAY(),${status_col}3<>"Done")'], fill=RED_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${status_col}3="Blocked"'], fill=AMBER_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${status_col}3="Done"'], font=GREY_FONT))
    ws.conditional_formatting.add(f"{upd_col}3:{upd_col}{LAST}", FormulaRule(
        formula=[f'AND(${upd_col}3<>"",${upd_col}3<TODAY()-14,${status_col}3<>"Done")'], font=ORANGE_FONT))

for i, owner in enumerate(OWNERS):
    ws = wb.create_sheet(owner); ws.sheet_properties.tabColor = TAB_COLORS[i]
    ws["A1"] = owner; ws["A1"].font = title_font
    ws["B1"] = f"Only {owner if owner != 'Team member 4' else 'this person'} and Suhail can edit this tab. One row per project; use the dropdowns."
    if owner == "Team member 4":
        ws["B1"] = "For the fourth team member: change the name in A1 and rename this tab. Only they and Suhail can edit it."
    ws["B1"].font = note_font
    ws.merge_cells("B1:K1")
    ws.row_dimensions[1].height = 26
    header(ws, 2, OWN_HDR)
    data = rows_by_owner.get(owner, [])
    if data:
        ws["A2"].comment = Comment("Source: Dabur Digital Network hub (claude.ai artifact), exported 7 Oct 2026.", "Tracker")
    for ri, row in enumerate(data, start=3):
        for ci, v in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=v)
    body(ws, f"A3:K{max(3, len(data) + 2)}")
    for col in ("F", "K"):
        for ri in range(3, 61):
            ws[f"{col}{ri}"].number_format = DATE_FMT
    add_validations(ws, {"B": ("Lists!$A$2:$A$30", False), "C": ("Lists!$B$2:$B$200", False),
                         "D": ("Lists!$C$2:$C$30", True), "E": ("Lists!$D$2:$D$10", True),
                         "I": ("Lists!$E$2:$E$60", False)})
    add_date_validation(ws, "F"); add_date_validation(ws, "K")
    row_colours(ws, "K", "E", "F", "K")
    widths(ws, OWN_W)
    ws.freeze_panes = "B3"

# ---------- All projects ----------
allp["A1"] = "All projects"; allp["A1"].font = title_font
allp["C1"] = "Fills itself in from each person's tab. Don't type here; change projects on the owner's tab."
allp["C1"].font = note_font
ALL_HDR = ["Owner"] + OWN_HDR
header(allp, 2, ALL_HDR)
parts = "; ".join(f"IF(LEN('{o}'!A3:A),'{o}'!$A$1,),'{o}'!A3:K" for o in OWNERS)
allp["A3"] = "=ARRAYFORMULA(QUERY({" + parts + "},\"select * where Col2 <> ''\",0))"
for col in ("G", "L"):
    for ri in range(3, 81):
        allp[f"{col}{ri}"].number_format = DATE_FMT
body(allp, "A3:L60")
row_colours(allp, "L", "F", "G", "L")
widths(allp, {"A": 13, "B": 38, "C": 15, "D": 20, "E": 22, "F": 13, "G": 12, "H": 48, "I": 32, "J": 22, "K": 26, "L": 13})
allp.freeze_panes = "C3"

# ---------- Comments ----------
cm = wb.create_sheet("Comments"); cm.sheet_properties.tabColor = "A26300"
cm["A1"] = "Comments"; cm["A1"].font = title_font
cm["C1"] = "Everyone can write here. Newest at the bottom; the Dashboard shows the latest 10. Suhail replies in column G."
cm["C1"].font = note_font
header(cm, 2, ["Date", "Who", "Project (or General)", "Type", "Comment", "Link", "Reply from Suhail"])
for ri in range(3, 61): cm[f"A{ri}"].number_format = DATE_FMT
add_date_validation(cm, "A")
for col, src, strict in (("B", "Lists!$F$2:$F$20", True), ("C", "'All projects'!$B$3:$B$500", False),
                         ("D", "Lists!$G$2:$G$20", True)):
    dv = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=strict)
    cm.add_data_validation(dv); dv.add(f"{col}3:{col}{LAST}")
cm.conditional_formatting.add(f"A3:G{LAST}", FormulaRule(formula=['$D3="Urgent flag"'], fill=RED_FILL))
cm.conditional_formatting.add(f"A3:G{LAST}", FormulaRule(formula=['$D3="Question for Suhail"'], fill=AMBER_FILL))
widths(cm, {"A": 12, "B": 13, "C": 34, "D": 18, "E": 60, "F": 28, "G": 40})
body(cm, "A3:G40")
cm.freeze_panes = "A3"

# ---------- Lists ----------
ls = wb.create_sheet("Lists"); ls.sheet_properties.tabColor = "5A6B62"
ls_cols = [("Markets", MARKET_LIST), ("Brands", settings.get("brands") or []), ("Streams", list(STREAMS.values())),
           ("Status", STATUS_LIST), ("Agencies", AGENCIES), ("People", PEOPLE_LIST), ("Comment types", TYPES)]
header(ls, 1, [h for h, _ in ls_cols])
for ci, (_, vals) in enumerate(ls_cols, start=1):
    for ri, v in enumerate(vals, start=2):
        c = ls.cell(row=ri, column=ci, value=v); c.font = body_font
widths(ls, {"A": 18, "B": 26, "C": 26, "D": 14, "E": 24, "F": 16, "G": 20})
ls.freeze_panes = "A2"

# ---------- Dashboard ----------
D = dash
D["B2"] = "Dashboard"; D["B2"].font = Font(name=FONT, size=18, bold=True, color=INK)
D["B3"] = "Fills itself in from everyone's tabs. Active means not done. Overdue means past the due date and not done."
D["B3"].font = note_font
A = "'All projects'!"
# by person
D["B5"] = "By person"; D["B5"].font = sec_font
p_hdr = ["Person", "Active", "In progress", "Blocked", "Overdue", "Due in next 14 days", "Done"]
header(D, 6, p_hdr, col=2)
r0 = 7
for i, o in enumerate(OWNERS):
    r = r0 + i
    D[f"B{r}"] = f"='{o}'!A1"
    D[f"C{r}"] = f'=COUNTIFS({A}$A$3:$A,$B{r},{A}$F$3:$F,"<>Done")'
    D[f"D{r}"] = f"=COUNTIFS({A}$A$3:$A,$B{r},{A}$F$3:$F,D$6)"
    D[f"E{r}"] = f"=COUNTIFS({A}$A$3:$A,$B{r},{A}$F$3:$F,E$6)"
    D[f"F{r}"] = f'=COUNTIFS({A}$A$3:$A,$B{r},{A}$F$3:$F,"<>Done",{A}$G$3:$G,"<"&TODAY())'
    D[f"G{r}"] = f'=COUNTIFS({A}$A$3:$A,$B{r},{A}$F$3:$F,"<>Done",{A}$G$3:$G,">="&TODAY(),{A}$G$3:$G,"<="&TODAY()+14)'
    D[f"H{r}"] = f"=COUNTIFS({A}$A$3:$A,$B{r},{A}$F$3:$F,H$6)"
rt = r0 + len(OWNERS)
D[f"B{rt}"] = "Team total"
for col in "CDEFGH":
    D[f"{col}{rt}"] = f"=SUM({col}{r0}:{col}{rt-1})"
body(D, f"B{r0}:H{rt}")
for col in "BCDEFGH":
    D[f"{col}{rt}"].font = Font(name=FONT, size=10, bold=True, color=INK)
    D[f"{col}{rt}"].border = Border(top=Side(style="thin", color=INK))
D.conditional_formatting.add(f"F{r0}:F{rt}", FormulaRule(formula=[f"F{r0}>0"], fill=RED_FILL))
D.conditional_formatting.add(f"E{r0}:E{rt}", FormulaRule(formula=[f"E{r0}>0"], fill=AMBER_FILL))

# by market
mr = rt + 2
D[f"B{mr}"] = "By market"; D[f"B{mr}"].font = sec_font
header(D, mr + 1, ["Market", "Active", "In progress", "Blocked", "Overdue", "Due in next 14 days", "Done"], col=2)
for i, m in enumerate(MARKET_LIST):
    r = mr + 2 + i
    D[f"B{r}"] = f"=Lists!A{2+i}"
    hr = mr + 1
    D[f"C{r}"] = f'=COUNTIFS({A}$C$3:$C,$B{r},{A}$F$3:$F,"<>Done")'
    D[f"D{r}"] = f"=COUNTIFS({A}$C$3:$C,$B{r},{A}$F$3:$F,D${hr})"
    D[f"E{r}"] = f"=COUNTIFS({A}$C$3:$C,$B{r},{A}$F$3:$F,E${hr})"
    D[f"F{r}"] = f'=COUNTIFS({A}$C$3:$C,$B{r},{A}$F$3:$F,"<>Done",{A}$G$3:$G,"<"&TODAY())'
    D[f"G{r}"] = f'=COUNTIFS({A}$C$3:$C,$B{r},{A}$F$3:$F,"<>Done",{A}$G$3:$G,">="&TODAY(),{A}$G$3:$G,"<="&TODAY()+14)'
    D[f"H{r}"] = f"=COUNTIFS({A}$C$3:$C,$B{r},{A}$F$3:$F,H${hr})"
mend = mr + 1 + len(MARKET_LIST)
body(D, f"B{mr+2}:H{mend}")
D.conditional_formatting.add(f"F{mr+2}:F{mend}", FormulaRule(formula=[f"F{mr+2}>0"], fill=RED_FILL))
D.conditional_formatting.add(f"E{mr+2}:E{mend}", FormulaRule(formula=[f"E{mr+2}>0"], fill=AMBER_FILL))

# needs attention
nr = mend + 2
D[f"B{nr}"] = "Needs attention: overdue or blocked"; D[f"B{nr}"].font = sec_font
header(D, nr + 1, ["Owner", "Project", "Market", "Status", "Due date", "Blocked on / need from Suhail"], col=2)
D[f"B{nr+2}"] = ("=IFERROR(QUERY(" + A + "A3:L,\"select A,B,C,F,G,I where F = 'Blocked' or (G < date '\"&TEXT(TODAY(),\"yyyy-mm-dd\")&\"' and F <> 'Done') order by G\",0),\"Nothing overdue or blocked.\")")
for ri in range(nr + 2, nr + 32): D[f"F{ri}"].number_format = DATE_FMT
body(D, f"B{nr+2}:G{nr+31}")

# due soon
sr = nr + 33
D[f"B{sr}"] = "Due in the next 14 days"; D[f"B{sr}"].font = sec_font
header(D, sr + 1, ["Owner", "Project", "Market", "Status", "Due date", "Latest update / next step"], col=2)
D[f"B{sr+2}"] = ("=IFERROR(QUERY(" + A + "A3:L,\"select A,B,C,F,G,H where F <> 'Done' and G >= date '\"&TEXT(TODAY(),\"yyyy-mm-dd\")&\"' and G <= date '\"&TEXT(TODAY()+14,\"yyyy-mm-dd\")&\"' order by G\",0),\"Nothing due in the next 14 days.\")")
for ri in range(sr + 2, sr + 32): D[f"F{ri}"].number_format = DATE_FMT
body(D, f"B{sr+2}:G{sr+31}")

# latest comments
cr = sr + 33
D[f"B{cr}"] = "Latest comments"; D[f"B{cr}"].font = sec_font
header(D, cr + 1, ["Date", "Who", "Project", "Type", "Comment", "Reply from Suhail"], col=2)
D[f"B{cr+2}"] = "=IFERROR(QUERY(Comments!A3:G,\"select A,B,C,D,E,G where E <> '' order by A desc limit 10\",0),\"No comments yet.\")"
for ri in range(cr + 2, cr + 13): D[f"B{ri}"].number_format = DATE_FMT
body(D, f"B{cr+2}:G{cr+12}")

widths(D, {"A": 3, "B": 16, "C": 38, "D": 15, "E": 14, "F": 14, "G": 44, "H": 10})
D.sheet_view.showGridLines = False

wb.save(OUT)
print("saved", OUT, {o: len(v) for o, v in rows_by_owner.items()})
